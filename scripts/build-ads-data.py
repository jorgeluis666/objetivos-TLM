#!/usr/bin/env python3
"""Genera data/tlm-ads-2026.json a partir de las exportaciones de Meta Ads de The Little Market.

Fuente: carpeta "Meta files - The Little Market" de Google Drive, una hoja por mes
("The Little Market - Octubre 2026"), pestana Raw Data Report: una fila por dia / edad / sexo / anuncio,
con los importes en soles (PEN).

Reglas:
  - El mes de cada hoja se reconoce por el nombre ("Octubre", "The Little Market - Octubre 2026") o por
    el rango de fechas de la exportacion ("..._20261001-20261031").
  - Solo cuentan las filas cuyo dia cae en ese mes. Si una hoja no trae ningun dia de su mes (por ejemplo,
    una copia del mes anterior), el mes queda sin datos y se avisa en "warnings".
  - Grupos de campana: Ventas (objetivo Ventas: compras en la web), Mensajes (WhatsApp y visitas al perfil
    de Instagram) y Reconocimiento (resultado ThruPlay).

Uso:
  python scripts/build-ads-data.py [--root "G:/Mi unidad/.../The Little Market"]
Requiere: openpyxl.
"""
import argparse
import calendar
import glob
import json
import os
import re
import sys
import unicodedata
import warnings
from datetime import date, datetime, timezone

import openpyxl

warnings.filterwarnings('ignore')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_DRIVE = 'G:/Mi unidad/Agencia Lima Retail/01_Clientes y Ventas/02_Clientes de la agencia/The Little Market'
META_DIR = 'Meta files - The Little Market'
OUT = os.path.join(ROOT, 'data', 'tlm-ads-2026.json')
YEAR = 2026
MONTHS = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
CURRENCY = 'PEN'
TOP_ADS = 5

# Encabezado de la exportacion (sin tildes, en minusculas) -> clave interna.
COLUMNS = {
    'dia': 'day', 'nombre del anuncio': 'ad', 'objetivo': 'objective', 'impresiones': 'impressions', 'alcance': 'reach',
    'clics en el enlace': 'clicks', 'visitas a la pagina de destino del sitio web': 'landingViews',
    'interacciones': 'interactions', 'articulos agregados al carrito': 'addToCart', 'pagos iniciados': 'checkouts',
    'compras': 'purchases', 'valor de conversion de compras': 'revenue',
    'conversaciones con mensajes iniciadas': 'conversations', 'tipo de resultado': 'resultType', 'resultados': 'results',
    'enlace de vista previa': 'preview', 'nombre de la campana': 'campaign',
}
SUMMED = ['spend', 'impressions', 'reach', 'clicks', 'landingViews', 'addToCart', 'checkouts', 'purchases', 'revenue',
          'conversations', 'interactions']
METRICS = SUMMED + ['profileVisits', 'thruplays']
MONEY = {'spend', 'revenue'}
AD_FIELDS = ['spend', 'impressions', 'clicks', 'purchases', 'revenue', 'conversations', 'profileVisits', 'thruplays']
# Orden de los mejores anuncios segun el objetivo de la campana.
TOP_SORT = {'Ventas': ('purchases', 'revenue', 'clicks'), 'Mensajes': ('conversations', 'profileVisits', 'clicks'),
            'Reconocimiento': ('thruplays', 'impressions')}
GROUPS = {'ventas': 'Ventas', 'reconocimiento': 'Reconocimiento', 'trafico': 'Tráfico', 'interaccion': 'Interacción',
          'clientes potenciales': 'Clientes potenciales'}


def norm(text):
    return unicodedata.normalize('NFD', str(text or '')).encode('ascii', 'ignore').decode().strip().lower()


def num(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if number != number else number


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.match(r'(\d{4})-(\d{2})-(\d{2})', str(value or ''))
    return date(int(match[1]), int(match[2]), int(match[3])) if match else None


def zero():
    return {key: 0.0 for key in METRICS}


def add(target, row):
    for key in SUMMED:
        target[key] += num(row.get(key))
    kind = norm(row.get('resultType'))
    # "Resultados" mezcla tipos: se separan los que se reportan aparte.
    if kind == 'thruplay':
        target['thruplays'] += num(row.get('results'))
    elif 'visitas al perfil' in kind:
        target['profileVisits'] += num(row.get('results'))


def rounded(values, keys=METRICS):
    return {key: round(values[key], 2) if key in MONEY else int(round(values[key])) for key in keys}


def group_of(name, objective):
    if 'mensaje' in norm(name) or 'mensaje' in norm(objective):
        return 'Mensajes'
    return GROUPS.get(norm(objective), str(objective or 'Otros').strip())


def month_file(folder, month):
    """Hoja del mes: por nombre ("Octubre.xlsx", "The Little Market - Octubre 2026.xlsx") o por el rango de
    fechas de la exportacion ("..._20261001-20261031.xlsx")."""
    number = MONTHS.index(month) + 1
    word = re.compile(r'(^|[^a-z])' + norm(month) + r'([^a-z]|$)')
    for path in sorted(glob.glob(os.path.join(folder, '*.xlsx'))):
        name = norm(os.path.basename(path))
        span = re.search(r'(\d{4})(\d{2})\d{2}\s*-\s*\d{8}', name)
        if word.search(name) or (span and int(span.group(1)) == YEAR and int(span.group(2)) == number):
            return path
    return None


def read_rows(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book['Raw Data Report'] if 'Raw Data Report' in book.sheetnames else book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    keys, currency = [], None
    for cell in next(rows, None) or []:
        spend = re.match(r'importe gastado \((\w+)\)', norm(cell))
        if spend:
            currency = spend.group(1).upper()
        keys.append('spend' if spend else COLUMNS.get(norm(cell)))
    if 'spend' not in keys or 'campaign' not in keys or 'day' not in keys:
        raise ValueError(f'{path}: no tiene las columnas de la exportacion de Meta (Dia, Importe gastado, Nombre de la campana)')
    if currency != CURRENCY:
        raise ValueError(f'{path}: moneda inesperada {currency} (se espera {CURRENCY})')
    records = []
    for values in rows:
        row = {key: value for key, value in zip(keys, values) if key}
        row['day'] = to_date(row.get('day'))
        row['campaign'] = str(row.get('campaign') or '').strip()
        if row['day'] and row['campaign']:
            records.append(row)
    book.close()
    return records


def top_ads(rows, group):
    ads = {}
    for row in rows:
        name = str(row.get('ad') or '').strip() or '(sin nombre)'
        ad = ads.setdefault(name, {'name': name, 'url': None, **zero()})
        add(ad, row)
        url = row.get('preview')
        if not ad['url'] and isinstance(url, str) and url.startswith('https://'):
            ad['url'] = url
    keys = TOP_SORT.get(group, ('clicks',))
    ranked = sorted((ad for ad in ads.values() if ad['spend'] > 0),
                    key=lambda ad: tuple(ad[key] for key in keys) + (ad['spend'],), reverse=True)[:TOP_ADS]
    result = []
    for ad in ranked:
        item = {'name': ad['name'], **rounded(ad, AD_FIELDS)}
        if ad['url']:
            item['url'] = ad['url']
        result.append(item)
    return result


def campaigns_of(rows):
    by_name = {}
    for row in rows:
        by_name.setdefault(row['campaign'], []).append(row)
    campaigns = []
    for name, items in by_name.items():
        totals = zero()
        for row in items:
            add(totals, row)
        if totals['spend'] <= 0:
            continue
        objective = next((str(row['objective']).strip() for row in items if row.get('objective')), '')
        group = group_of(name, objective)
        campaign = {'name': name, 'objective': objective, 'group': group, **rounded(totals)}
        campaign['cpm'] = round(totals['spend'] / totals['impressions'] * 1000, 2) if totals['impressions'] else None
        campaign['cpc'] = round(totals['spend'] / totals['clicks'], 4) if totals['clicks'] else None
        campaign['topAds'] = top_ads(items, group)
        campaigns.append(campaign)
    order = {'Ventas': 0, 'Mensajes': 1, 'Reconocimiento': 2}
    return sorted(campaigns, key=lambda c: (order.get(c['group'], 3), -c['spend']))


def sheet_ids(drive):
    """IDs de Drive de cada hoja, tomados del manifest que deja drive-download.py (si existe)."""
    try:
        with open(os.path.join(drive, 'manifest.json'), encoding='utf-8') as handle:
            folder = json.load(handle).get(META_DIR, {})
    except FileNotFoundError:
        return None, {}
    return folder.get('folderId'), {item['file']: item['id'] for item in folder.get('files', []) if item.get('file')}


def short(day):
    return f'{day.day}/{day.month:02d}'


def build(drive):
    folder = os.path.join(drive, META_DIR)
    if not os.path.isdir(folder):
        sys.exit(f'No existe la carpeta {folder}')
    folder_id, ids = sheet_ids(drive)
    months, warnings_list = [], []
    for index, month in enumerate(MONTHS):
        path = month_file(folder, month)
        entry = {'name': month, 'status': 'sin-datos', 'period': None, 'sheet': None, 'firstDay': None, 'lastDay': None,
                 'kpis': None, 'spendByGroup': {}, 'campaigns': []}
        months.append(entry)
        if not path:
            continue
        file_name = os.path.basename(path)
        entry['sheet'] = {'title': os.path.splitext(file_name)[0], 'id': ids.get(file_name)}
        rows = read_rows(path)
        inside = [row for row in rows if row['day'].year == YEAR and row['day'].month == index + 1]
        if rows and not inside:
            days = sorted(row['day'] for row in rows)
            warnings_list.append(f'{month}: la hoja "{entry["sheet"]["title"]}" no trae dias de {month.lower()} '
                                 f'(sus fechas van del {short(days[0])} al {short(days[-1])}). Revisa que no sea una copia de otro mes.')
            entry['notice'] = 'La hoja no trae dias de este mes'
            continue
        campaigns = campaigns_of(inside)
        if not campaigns:
            continue
        totals = zero()
        for row in inside:
            add(totals, row)
        days = sorted(row['day'] for row in inside if num(row.get('spend')) > 0)
        by_group = {}
        for campaign in campaigns:
            by_group[campaign['group']] = round(by_group.get(campaign['group'], 0) + campaign['spend'], 2)
        entry.update({'firstDay': days[0].isoformat(), 'lastDay': days[-1].isoformat(), 'kpis': rounded(totals),
                      'spendByGroup': by_group, 'campaigns': campaigns})

    with_data = [i for i, month in enumerate(months) if month['kpis']]
    cutoff = None
    for index in with_data:
        month = months[index]
        days_in_month = calendar.monthrange(YEAR, index + 1)[1]
        last = int(month['lastDay'][8:10])
        # Un mes con datos posteriores esta cerrado; el ultimo, si llega al penultimo dia (la pauta pudo terminar antes).
        closed = index < with_data[-1] or last >= days_in_month - 1
        end = days_in_month if closed else last
        month['status'] = 'cerrado' if closed else 'parcial'
        month['period'] = f'1 - {end} {month["name"]} {YEAR}'
        cutoff = date(YEAR, index + 1, end).isoformat()
    for index in range(with_data[-1] if with_data else 0):
        month = months[index]
        if not month['sheet']:
            warnings_list.append(f'{month["name"]}: no hay hoja del mes en la carpeta de Drive.')
            month['notice'] = 'Sin hoja en la carpeta de Drive'
    return {
        'account': 'The Little Market',
        'year': YEAR,
        'cutoff': cutoff,
        'generatedAt': date.today().isoformat(),
        # El boton "Sincronizar con Drive" del tablero espera a que cambie este sello para recargar.
        'syncedAt': datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
        'source': 'Exportaciones de Meta Ads (Google Drive)',
        'platform': {'key': 'meta', 'label': 'Meta Ads', 'currency': CURRENCY, 'symbol': 'S/'},
        'folderId': folder_id or os.environ.get('TLM_META_FOLDER_ID', '1lpi0wzaLuwL3MWMKBNtHBHJyDBD6p8Ie'),
        'warnings': warnings_list,
        'months': months,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--root', default=os.environ.get('TLM_DRIVE_ROOT', DEFAULT_DRIVE), help='Carpeta The Little Market en Google Drive')
    args = parser.parse_args()
    if not os.path.isdir(args.root):
        sys.exit(f'No existe la carpeta {args.root}')
    data = build(args.root)
    with open(OUT, 'w', encoding='utf-8', newline='\n') as handle:
        json.dump(data, handle, ensure_ascii=False, indent=1)
        handle.write('\n')
    for month in data['months']:
        if not month['kpis']:
            continue
        k = month['kpis']
        roas = k['revenue'] / k['spend'] if k['spend'] else 0
        print(f"{month['name']:<11} {month['status']:<8} S/ {k['spend']:>9,.2f} | {k['purchases']:>3} compras "
              f"S/ {k['revenue']:>9,.2f} (ROAS {roas:.2f}) | {k['conversations']:>4} conversaciones | {len(month['campaigns'])} campanas")
    for warning in data['warnings']:
        print(f'[aviso] {warning}')
    print(f'[ads] escrito {os.path.relpath(OUT, ROOT)} (corte {data["cutoff"]})')


if __name__ == '__main__':
    main()
